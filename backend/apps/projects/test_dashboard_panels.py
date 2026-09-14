"""The dashboard workbook's summary panels — duration, submittal grids and
financial progress by BOQ (register F3-F5).

The sheet below copies the airport dashboard's own layout cell for cell where
it matters: the duration and time-performance panels side by side sharing the
"Remaining Duration" wording, the shop-drawing and material grids a few
columns apart on one row, the empty "previous" grids beneath them, and the BOQ
panel's category row carrying the contract total at its start.
"""
import io

import openpyxl
from django.test import SimpleTestCase, TestCase

from apps.accounts.models import Company

from .dashboard_panels import (import_dashboard_panels, parse_boq, parse_dashboard_panels,
                               parse_duration, parse_submittals)
from .models import DashboardPanels, Project


def _cell(ref):
    """(row, column), 0-based, from an "AB12" reference."""
    letters = "".join(ch for ch in ref if ch.isalpha())
    col = 0
    for ch in letters:
        col = col * 26 + ord(ch.upper()) - 64
    return int(ref[len(letters):]) - 1, col - 1


AIRPORT = {
    "V15": "PROJECT DURATION", "W15": 1380, "X15": "Time Performance", "Z15": 0,
    "AA15": "PROJECT DURATION", "AB15": 1380,
    "V16": "COMPLETED DURATION", "W16": 1313, "X16": "Ellapsed Duration", "Y16": 0.9514492753623188,
    "AA16": "DELAY", "AB16": -47,
    "V17": "REMAINING DURATION", "W17": 67, "X17": "Remaining Duration ", "Y17": 0.04855072463768116,
    "A20": "PROJECT DELAY IN CALENDAR DAYS ", "C20": -47,

    "Z82": "SHOP DRAWINGS ", "AA82": "SUBMITTED", "AB82": "APPROVED", "AC82": "REJECTED", "AD82": "pending",
    "AF82": "MATERIAL SUBMITTALS", "AG82": "SUBMITTED", "AH82": "APPROVED", "AI82": "REJECTED", "AJ82": "pending",
    "Z83": "CIVIL", "AA83": 426, "AB83": 424, "AC83": 2, "AD83": 0,
    "AF83": "CIVIL", "AG83": 165, "AH83": 165, "AI83": 0, "AJ83": 0,
    "Z84": "MECHANICAL", "AA84": 82, "AB84": 82, "AC84": 0, "AD84": 0,
    "AF84": "Arch ", "AG84": 94, "AH84": 93, "AI84": 1, "AJ84": 0,
    "Z85": "ELECTRICAL", "AA85": 100, "AB85": 100, "AC85": 0, "AD85": 0,
    "AF85": "MEP ", "AG85": 71, "AH85": 70, "AI85": 1, "AJ85": 0,
    # The previous month's grids: headed the same, holding no counts.
    "Z91": "SHOP DRAWINGS ", "AA91": "SUBMITTED", "AB91": "APPROVED", "AC91": "REJECTED", "AD91": "PENDING",
    "AF91": "MATERIAL SUBMITTALS", "AG91": "SUBMITTED", "AH91": "APPROVED", "AI91": "REJECTED", "AJ91": "pending",
    "Z92": "ARCH", "AF92": "ARCH", "Z93": "CIVIL", "AF93": "CIVIL",

    "AB116": "Financial Progress according to BOQ",
    "V117": 685661117.68, "W117": "اعمال الاعتيادي", "X117": "اعمال الكهرباء",
    "Z117": "اعمال الميكانيكا", "AA117": "اعمال الجداريات", "AB117": "اعمال الفرش",
    "AA118": 33689800.36, "AB118": 55603647,
    "V119": "Budget ", "W119": 0, "X119": 0, "Z119": 0, "AA119": 0.0491347686069653, "AB119": 0.08109493971036343,
    "V121": "Actual", "W121": 0.2892621879268702, "X121": 0.1492665844864421,
    "Z121": 0.032327989118372265, "AA121": 0.044372584987142624, "AB121": 0,
}


def _rows(cells=AIRPORT):
    placed = [(_cell(ref), value) for ref, value in cells.items()]
    height = max(r for (r, _), _ in placed) + 1
    width = max(c for (_, c), _ in placed) + 1
    rows = [[None] * width for _ in range(height)]
    for (r, c), value in placed:
        rows[r][c] = value
    return rows


def _workbook(cells=AIRPORT, title="Dashboard"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for ref, value in cells.items():
        ws[ref] = value
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class DurationPanelTests(SimpleTestCase):
    def test_reads_the_duration_block_and_its_delay(self):
        self.assertEqual(parse_duration(_rows()), {
            "project_days": 1380.0, "completed_days": 1313.0, "remaining_days": 67.0,
            "elapsed_pct": 0.9514492753623188, "remaining_pct": 0.04855072463768116,
            "delay_days": -47.0,
        })

    def test_remaining_days_and_remaining_share_are_told_apart_by_value(self):
        """Both panels say "Remaining Duration"; a day count is never a
        fraction, so each lands in its own key whichever is read first."""
        swapped = {"X17": "REMAINING DURATION", "Y17": 67,
                   "V17": "Remaining Duration ", "W17": 0.04855072463768116,
                   "V15": "PROJECT DURATION", "W15": 1380}
        out = parse_duration(_rows(swapped))
        self.assertEqual(out["remaining_days"], 67.0)
        self.assertEqual(out["remaining_pct"], 0.04855072463768116)

    def test_a_sheet_without_the_block_reads_nothing(self):
        self.assertEqual(parse_duration(_rows({"A1": "Notes", "B1": 3})), {})


class SubmittalPanelTests(SimpleTestCase):
    def test_each_grid_reads_its_own_status_columns(self):
        """The grids share a row a few columns apart. Scanning a fixed width
        took the material grid's SUBMITTED for the shop drawings and reported
        165 civil drawings submitted where the sheet says 426."""
        grids = parse_submittals(_rows())
        self.assertEqual(grids["shop_drawing"], [
            {"discipline": "CIVIL", "submitted": 426, "approved": 424, "rejected": 2, "pending": 0},
            {"discipline": "MECHANICAL", "submitted": 82, "approved": 82, "rejected": 0, "pending": 0},
            {"discipline": "ELECTRICAL", "submitted": 100, "approved": 100, "rejected": 0, "pending": 0},
        ])
        self.assertEqual(grids["material"], [
            {"discipline": "CIVIL", "submitted": 165, "approved": 165, "rejected": 0, "pending": 0},
            {"discipline": "Arch", "submitted": 94, "approved": 93, "rejected": 1, "pending": 0},
            {"discipline": "MEP", "submitted": 71, "approved": 70, "rejected": 1, "pending": 0},
        ])

    def test_the_empty_previous_grid_is_not_taken_for_the_current_one(self):
        current_rows = range(81, 85)          # sheet rows 82-85, 0-based
        only_previous = {k: v for k, v in AIRPORT.items() if _cell(k)[0] not in current_rows}
        self.assertEqual(parse_submittals(_rows(only_previous)), {})


class BoqPanelTests(SimpleTestCase):
    def test_reads_categories_with_their_budget_and_actual_shares(self):
        panel = parse_boq(_rows())
        self.assertEqual(panel["total"], 685661117.68)
        self.assertEqual([r["category"] for r in panel["rows"]],
                         ["اعمال الاعتيادي", "اعمال الكهرباء", "اعمال الميكانيكا",
                          "اعمال الجداريات", "اعمال الفرش"])
        civil = panel["rows"][0]
        self.assertEqual((civil["budget_share"], civil["financial_percent"]), (0.0, 0.2892621879268702))
        self.assertEqual(panel["rows"][4]["budget_share"], 0.08109493971036343)

    def test_no_panel_reads_nothing(self):
        self.assertEqual(parse_boq(_rows({"A1": "Budget", "B1": 1})), {})


class DashboardPanelsImportTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name="Acme")
        self.project = Project.objects.create(
            company=self.company, name="Tower", project_type="commercial")

    def test_parses_every_panel_off_the_dashboard_sheet_only(self):
        wb = openpyxl.load_workbook(_workbook(), data_only=True)
        self.assertEqual(sorted(parse_dashboard_panels(wb)), ["boq", "duration", "submittals"])
        wb = openpyxl.load_workbook(_workbook(title="notes"), data_only=True)
        self.assertEqual(parse_dashboard_panels(wb), {})

    def test_an_import_replaces_the_stored_panels(self):
        summary = import_dashboard_panels(self.project, _workbook())
        self.assertEqual(summary, {"duration": 6, "submittals": 6, "boq": 5})
        import_dashboard_panels(self.project, _workbook({"V15": "PROJECT DURATION", "W15": 900}))
        stored = DashboardPanels.objects.get(project=self.project)
        self.assertEqual(stored.data, {"duration": {"project_days": 900.0}})
        self.assertEqual(stored.company, self.company)

    def test_a_workbook_without_panels_leaves_the_stored_ones_alone(self):
        import_dashboard_panels(self.project, _workbook())
        with self.assertRaises(ValueError):
            import_dashboard_panels(self.project, _workbook(title="notes"))
        self.assertIn("boq", DashboardPanels.objects.get(project=self.project).data)
